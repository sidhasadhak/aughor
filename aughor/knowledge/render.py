"""Markdown → any format the platform can render.

The other half of the pivot. `knowledge.convert` brings fourteen formats IN by
turning them into Markdown; this takes that Markdown back OUT. Keeping one canonical
form in the middle is what makes "any format to any format" fourteen readers plus six
writers instead of eighty-four converters.

The writers were already here
-----------------------------
`aughor/export/` has rendered PDFs and PowerPoint decks for a long time, driven by an
`ExportDoc` — a list of typed blocks (heading, prose, bullets, table, code). It was
wired to exactly one producer, `report_json` from an investigation, so the renderers
could only ever emit reports. Nothing about them was report-specific; they just had a
single caller.

So the only genuinely new thing here is a Markdown parser that produces `ExportDoc`.
Point it at the renderers and PDF and PPTX arrive for free, correctly styled, with the
table handling and the pagination someone already got right.

Why the parser is written rather than pulled in
-----------------------------------------------
Its input is not arbitrary Markdown off the internet — it is anydoc's output, which is
well-formed GFM: ATX headings, pipe tables, fenced code, `-` bullets. A dependency
would buy tolerance this input does not need, and the two output paths that matter for
safety (HTML, and text inside a PDF) want structure, not passthrough. Parsing to
blocks and building HTML from those blocks means every string is escaped on the way
out and no source document can inject markup into a page. A Markdown-to-HTML library
would let raw HTML in the source through by default, which for uploaded files is the
wrong default.

Six formats, one model. `md` is the pivot returned as-is; `txt` is it with the syntax
taken off; `html`, `pdf`, `docx` and `pptx` are rendered from the blocks.
"""
from __future__ import annotations

import html as _html
import io
import logging
import re
from dataclasses import dataclass

logger = logging.getLogger(__name__)

#: What a person can ask a document to become, with the MIME type and extension each
#: implies. Also the ORDER they are offered in: the pivot first, then the formats you
#: read, then the ones you edit.
FORMATS: dict[str, tuple[str, str, str]] = {
    "md":   ("Markdown",   "text/markdown; charset=utf-8", ".md"),
    "txt":  ("Plain text", "text/plain; charset=utf-8",    ".txt"),
    "html": ("HTML",       "text/html; charset=utf-8",     ".html"),
    "pdf":  ("PDF",        "application/pdf",              ".pdf"),
    "docx": ("Word",
             "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
             ".docx"),
    "pptx": ("PowerPoint",
             "application/vnd.openxmlformats-officedocument.presentationml.presentation",
             ".pptx"),
}

#: Formats needing the `export` extra (reportlab / python-pptx). `docx` is not among
#: them — python-docx rides the `docs` extra that read the file in the first place.
_NEEDS_EXPORT_EXTRA = {"pdf", "pptx"}


class RenderError(Exception):
    """A document could not be rendered into the requested format."""

    def __init__(self, message: str, *, code: str = "render_failed") -> None:
        super().__init__(message)
        self.code = code


# ── Markdown → blocks ─────────────────────────────────────────────────────────

_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_ORDERED = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_FENCE = re.compile(r"^\s*```+\s*(\w*)\s*$")
_TABLE_ROW = re.compile(r"^\s*\|(.+)\|\s*$")
_TABLE_RULE = re.compile(r"^\s*\|[\s:|-]+\|\s*$")
#: Inline syntax, stripped for the plain-text and slide paths. Images before links,
#: because `![alt](src)` is a link pattern with a `!` in front of it.
_INLINE = [
    (re.compile(r"!\[([^\]]*)\]\([^)]*\)"), r"\1"),   # image → its alt text
    (re.compile(r"\[([^\]]*)\]\([^)]*\)"), r"\1"),    # link  → its label
    (re.compile(r"\*\*([^*]+)\*\*"), r"\1"),
    (re.compile(r"(?<!\*)\*([^*]+)\*(?!\*)"), r"\1"),
    (re.compile(r"`([^`]+)`"), r"\1"),
]


def strip_inline(text: str) -> str:
    """Markdown emphasis and link syntax removed, the words kept."""
    for pattern, replacement in _INLINE:
        text = pattern.sub(replacement, text)
    return text.strip()


def _split_row(line: str) -> list[str]:
    """A pipe-table row's cells. The outer pipes are delimiters, not content.

    Strips exactly ONE pipe from each end, not every pipe there. `str.strip("|")` is
    greedy, so a row that opens with an empty leading cell — `||Highlights|`, which is
    how a converted slide renders a blank corner — had BOTH pipes eaten and came back
    one cell narrower than the rows beneath it. That mismatch is what fed the
    truncation below, and between them they deleted a table's entire contents.
    """
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    return [c.strip() for c in text.split("|")]


@dataclass
class _Table:
    columns: list[str]
    rows: list[list[str]]


def parse_markdown(markdown: str, title: str = "Document"):
    """Markdown → `ExportDoc`, the shape every renderer here already understands.

    Imported lazily: `ExportDoc` lives behind the `export` extra's import guard, and
    the `md`/`txt` paths must keep working on an install that has no renderers at all.
    """
    from aughor.export.document import Block, ExportDoc

    blocks: list[Block] = []
    paragraph: list[str] = []
    bullets: list[str] = []

    def flush_paragraph() -> None:
        if paragraph:
            blocks.append(Block("prose", text=strip_inline(" ".join(paragraph))))
            paragraph.clear()

    def flush_bullets() -> None:
        if bullets:
            blocks.append(Block("bullets", items=[strip_inline(b) for b in bullets]))
            bullets.clear()

    def flush() -> None:
        flush_paragraph()
        flush_bullets()

    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]

        fence = _FENCE.match(line)
        if fence:
            flush()
            lang = fence.group(1)
            body: list[str] = []
            i += 1
            while i < len(lines) and not _FENCE.match(lines[i]):
                body.append(lines[i])
                i += 1
            i += 1                                   # step over the closing fence
            blocks.append(Block("code", text="\n".join(body), caption=lang))
            continue

        # A pipe table needs its `| --- |` rule on the NEXT line; without it a line of
        # pipes is prose that happens to contain them.
        if (_TABLE_ROW.match(line) and i + 1 < len(lines)
                and _TABLE_RULE.match(lines[i + 1])):
            flush()
            columns = [strip_inline(c) for c in _split_row(line)]
            i += 2
            rows: list[list[str]] = []
            while i < len(lines) and _TABLE_ROW.match(lines[i]):
                rows.append([strip_inline(c) for c in _split_row(lines[i])])
                i += 1
            # The table is as wide as its WIDEST row, never as wide as its header.
            #
            # This used to truncate to the header's width, and that silently deleted
            # data on real input. Converters do not escape pipes inside a cell, so a
            # slide whose row label is "Luxury | Mytheresa" emits three cells under a
            # two-cell header — and truncation kept "Luxury", dropped "Mytheresa", and
            # threw away the entire bullet list of results beside it. The rendered PDF
            # looked perfectly well-formed with its numbers gone, which is the worst
            # shape a failure can take.
            #
            # Widening can leave an empty column; truncating loses content. Only one of
            # those is recoverable by the person looking at it.
            width = max([len(columns)] + [len(r) for r in rows]) if rows else len(columns)
            columns += [""] * (width - len(columns))
            rows = [r + [""] * (width - len(r)) for r in rows]
            blocks.append(Block("table", columns=columns, rows=rows))
            continue

        heading = _HEADING.match(line)
        if heading:
            flush()
            blocks.append(Block("heading", text=strip_inline(heading.group(2))))
            i += 1
            continue

        bullet = _BULLET.match(line) or _ORDERED.match(line)
        if bullet:
            flush_paragraph()
            bullets.append(bullet.group(1))
            i += 1
            continue

        if not line.strip():
            flush()
            i += 1
            continue

        flush_bullets()
        paragraph.append(line.strip())
        i += 1

    flush()
    return ExportDoc(title=title, kind="document", blocks=blocks)


# ── Renderers ─────────────────────────────────────────────────────────────────

def _to_text(doc) -> bytes:
    """Plain text: the structure kept as layout, the syntax gone.

    Tables become aligned columns rather than pipes, because the point of asking for
    text is that it reads without a renderer.
    """
    out: list[str] = []
    for block in doc.blocks:
        if block.kind == "heading":
            out += [block.text, "=" * len(block.text), ""]
        elif block.kind == "prose":
            out += [block.text, ""]
        elif block.kind == "bullets":
            out += [f"  • {item}" for item in block.items] + [""]
        elif block.kind == "code":
            out += [f"    {line}" for line in block.text.splitlines()] + [""]
        elif block.kind == "table":
            grid = [block.columns] + [[str(c) for c in row] for row in block.rows]
            widths = [max(len(r[i]) for r in grid) for i in range(len(block.columns))]
            for index, row in enumerate(grid):
                out.append("  ".join(cell.ljust(widths[i]) for i, cell in enumerate(row)).rstrip())
                if index == 0:
                    out.append("  ".join("-" * w for w in widths))
            out.append("")
    return "\n".join(out).encode("utf-8")


def _to_html(doc) -> bytes:
    """A standalone HTML page, with every string escaped on the way out.

    Built from the blocks rather than by converting Markdown directly, so nothing in
    an uploaded document can inject markup into the page. That is not a theoretical
    concern here: these documents arrive from outside and are rendered for other
    people to look at.
    """
    def esc(value) -> str:
        return _html.escape(str(value), quote=True)

    body: list[str] = []
    for block in doc.blocks:
        if block.kind == "heading":
            body.append(f"<h2>{esc(block.text)}</h2>")
        elif block.kind == "prose":
            body.append(f"<p>{esc(block.text)}</p>")
        elif block.kind == "bullets":
            items = "".join(f"<li>{esc(i)}</li>" for i in block.items)
            body.append(f"<ul>{items}</ul>")
        elif block.kind == "code":
            body.append(f"<pre><code>{esc(block.text)}</code></pre>")
        elif block.kind == "table":
            head = "".join(f"<th>{esc(c)}</th>" for c in block.columns)
            rows = "".join(
                "<tr>" + "".join(f"<td>{esc(c)}</td>" for c in row) + "</tr>"
                for row in block.rows)
            body.append(f"<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table>")

    return f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(doc.title)}</title>
<style>
  :root {{ color-scheme: light dark; }}
  body {{ margin: 0 auto; padding: 2rem 1.25rem; max-width: 46rem; line-height: 1.6;
         font: 16px/1.6 system-ui, -apple-system, "Segoe UI", sans-serif; }}
  h1 {{ font-size: 1.6rem; margin-bottom: 1.5rem; }}
  h2 {{ font-size: 1.2rem; margin-top: 2rem; }}
  table {{ border-collapse: collapse; width: 100%; margin: 1rem 0; display: block;
          overflow-x: auto; }}
  th, td {{ border: 1px solid #8883; padding: .4rem .6rem; text-align: left; }}
  th {{ background: #8881; }}
  pre {{ background: #8881; padding: .75rem; border-radius: 4px; overflow-x: auto; }}
</style></head>
<body><h1>{esc(doc.title)}</h1>
{chr(10).join(body)}
</body></html>""".encode("utf-8")


def _to_docx(doc) -> bytes:
    """A Word document, via python-docx — which was already here to READ .docx files."""
    try:
        import docx
    except ImportError:
        raise RenderError("Word output needs python-docx. Install with: "
                          "uv pip install -e '.[docs]'", code="renderer_missing")

    out = docx.Document()
    out.add_heading(doc.title, level=0)
    for block in doc.blocks:
        if block.kind == "heading":
            out.add_heading(block.text, level=1)
        elif block.kind == "prose":
            out.add_paragraph(block.text)
        elif block.kind == "bullets":
            for item in block.items:
                out.add_paragraph(item, style="List Bullet")
        elif block.kind == "code":
            out.add_paragraph(block.text, style="No Spacing")
        elif block.kind == "table":
            if not block.columns:
                continue
            table = out.add_table(rows=1, cols=len(block.columns))
            table.style = "Table Grid"
            for i, column in enumerate(block.columns):
                table.rows[0].cells[i].text = str(column)
            for row in block.rows:
                cells = table.add_row().cells
                for i, value in enumerate(row[:len(block.columns)]):
                    cells[i].text = str(value)
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def render(markdown: str, fmt: str, title: str = "Document") -> tuple[bytes, str, str]:
    """Render Markdown into `fmt`. Returns (bytes, media_type, suffix)."""
    key = (fmt or "").strip().lower().lstrip(".")
    if key not in FORMATS:
        raise RenderError(
            f"Unknown format '{fmt}'. Available: {', '.join(FORMATS)}.",
            code="unknown_format")
    _, media_type, suffix = FORMATS[key]

    if key == "md":
        return markdown.encode("utf-8"), media_type, suffix

    if key in _NEEDS_EXPORT_EXTRA:
        from aughor.export import EXPORT_AVAILABLE
        if not EXPORT_AVAILABLE:
            raise RenderError(
                f"{FORMATS[key][0]} output needs the 'export' extra. Install with: "
                f"uv pip install -e '.[export]'", code="renderer_missing")

    doc = parse_markdown(markdown, title=title)

    try:
        if key == "txt":
            return _to_text(doc), media_type, suffix
        if key == "html":
            return _to_html(doc), media_type, suffix
        if key == "docx":
            return _to_docx(doc), media_type, suffix
        if key == "pdf":
            from aughor.export import render_pdf
            return render_pdf(doc), media_type, suffix
        if key == "pptx":
            from aughor.export import render_pptx
            return render_pptx(doc), media_type, suffix
    except RenderError:
        raise
    except Exception as exc:
        logger.exception("Rendering to %s failed", key)
        raise RenderError(f"Could not render this document as {FORMATS[key][0]}: {exc}")

    raise RenderError(f"Unhandled format '{key}'.", code="unknown_format")  # pragma: no cover


def available_formats() -> list[dict]:
    """The formats this deployment can actually produce, and which need an extra.

    Reported rather than assumed, so the UI offers PDF only where reportlab is
    installed instead of offering it everywhere and failing at the click.
    """
    from aughor.export import EXPORT_AVAILABLE

    try:
        import docx  # noqa: F401
        has_docx = True
    except ImportError:
        has_docx = False

    out = []
    for key, (label, media_type, suffix) in FORMATS.items():
        if key in _NEEDS_EXPORT_EXTRA:
            ready = EXPORT_AVAILABLE
        elif key == "docx":
            ready = has_docx
        else:
            ready = True
        out.append({"format": key, "label": label, "media_type": media_type,
                    "suffix": suffix, "available": ready})
    return out
